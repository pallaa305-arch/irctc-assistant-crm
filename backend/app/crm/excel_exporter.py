import os
import shutil
from pathlib import Path
from datetime import datetime
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pandas as pd
from app.config import settings, BOOKINGS_DIR, EXPORTS_DIR, BACKUPS_DIR

EXCEL_COLUMNS = [
    "Booking ID",
    "PNR",
    "Transaction ID",
    "Booking Date",
    "Journey Date",
    "From",
    "To",
    "Boarding Station",
    "Train Number",
    "Train Name",
    "Departure",
    "Arrival",
    "Class",
    "Quota",
    "Passenger Count",
    "Passenger Names",
    "Booking Status",
    "Fare",
    "Payment Status",
    "Telegram Status",
    "WhatsApp Status",
    "Notes"
]

def init_excel_workbook(file_path: Optional[str] = None) -> openpyxl.Workbook:
    """Initializes a beautifully styled Excel workbook if it doesn't already exist."""
    path = file_path or settings.EXCEL_FILE_PATH
    if os.path.exists(path):
        return openpyxl.load_workbook(path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IRCTC Bookings"

    # Header styling (Green theme)
    header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    ws.append(EXCEL_COLUMNS)
    for col_num, col_name in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        # Auto-adjust column width
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = max(len(col_name) + 4, 14)

    wb.save(path)
    return wb

def append_booking_to_excel(booking, passengers: list, file_path: Optional[str] = None):
    """
    Appends a new booking row to the Excel file WITHOUT overwriting existing rows.
    """
    path = file_path or settings.EXCEL_FILE_PATH
    wb = init_excel_workbook(path)
    ws = wb["IRCTC Bookings"]

    passenger_names = ", ".join([p.name for p in passengers]) if passengers else ""

    row = [
        booking.booking_ref,
        booking.pnr or "N/A",
        booking.transaction_id or "N/A",
        booking.created_at.strftime("%d/%m/%Y %H:%M"),
        booking.journey_date.strftime("%d/%m/%Y"),
        booking.from_station,
        booking.to_station,
        booking.boarding_station or booking.from_station,
        booking.train_number or "N/A",
        booking.train_name or "N/A",
        booking.departure_time or "N/A",
        booking.arrival_time or "N/A",
        booking.journey_class,
        booking.quota,
        booking.passenger_count,
        passenger_names,
        booking.status,
        booking.fare or 0.0,
        booking.payment_status,
        booking.telegram_status,
        booking.whatsapp_status,
        booking.notes or ""
    ]

    ws.append(row)

    # Style appended row
    row_idx = ws.max_row
    font = Font(name="Calibri", size=10)
    for col_idx in range(1, len(EXCEL_COLUMNS) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = font
        cell.alignment = Alignment(vertical="center")

    wb.save(path)

def export_bookings_csv(bookings_list: List[dict], output_filename: Optional[str] = None) -> str:
    """Exports a list of bookings to CSV."""
    df = pd.DataFrame(bookings_list)
    filename = output_filename or f"bookings_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    out_path = EXPORTS_DIR / filename
    df.to_csv(out_path, index=False, encoding="utf-8-sig")
    return str(out_path)

def backup_database() -> str:
    """Creates a timestamped snapshot of the SQLite database."""
    db_file = Path(settings.DATABASE_URL.replace("sqlite:///", ""))
    if not db_file.exists():
        raise FileNotFoundError("Database file does not exist yet.")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file = BACKUPS_DIR / f"assistant_backup_{timestamp}.db"
    shutil.copy2(db_file, backup_file)
    return str(backup_file)
