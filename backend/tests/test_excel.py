import os
from datetime import datetime, date
from app.crm.excel_exporter import init_excel_workbook, append_booking_to_excel
import openpyxl

class DummyBooking:
    def __init__(self):
        self.booking_ref = "BK-TEST-12345"
        self.pnr = "2458917263"
        self.transaction_id = "TXN998877"
        self.created_at = datetime(2026, 9, 19, 10, 30)
        self.journey_date = date(2026, 9, 25)
        self.from_station = "NDLS"
        self.to_station = "BPL"
        self.boarding_station = "NDLS"
        self.train_number = "12002"
        self.train_name = "Bhopal Shatabdi"
        self.departure_time = "06:00"
        self.arrival_time = "14:40"
        self.journey_class = "CC"
        self.quota = "GN"
        self.passenger_count = 1
        self.fare = 1250.0
        self.status = "CONFIRMED"
        self.payment_status = "COMPLETED"
        self.telegram_status = "SENT"
        self.whatsapp_status = "NOT_SENT"
        self.notes = "Test dummy note"

class DummyPassenger:
    def __init__(self, name="John Doe"):
        self.name = name

def test_excel_init_and_append(tmp_path):
    test_excel = str(tmp_path / "test_bookings.xlsx")
    wb = init_excel_workbook(test_excel)
    assert os.path.exists(test_excel)

    # Append first row
    booking = DummyBooking()
    passengers = [DummyPassenger("John Doe")]
    append_booking_to_excel(booking, passengers, test_excel)

    # Append second row (Verifying non-destructive append)
    booking2 = DummyBooking()
    booking2.booking_ref = "BK-TEST-67890"
    booking2.pnr = "9876543210"
    append_booking_to_excel(booking2, passengers, test_excel)

    loaded_wb = openpyxl.load_workbook(test_excel)
    ws = loaded_wb["IRCTC Bookings"]
    assert ws.max_row == 3  # Header + 2 rows
    assert ws.cell(row=2, column=1).value == "BK-TEST-12345"
    assert ws.cell(row=3, column=1).value == "BK-TEST-67890"
